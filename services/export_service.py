"""
services/export_service.py
-------------------------------
엑셀/PDF 내보내기 유틸리티. SRS FR-ORD-05(주문 데이터 엑셀 내보내기),
FR-REPORT-04(보고서 엑셀/PDF 다운로드 - 주문/교환/반품/취소 내역,
월별손익보고서, 광고성과보고서) 대응.

계산/상태 변경이 없는 단순 직렬화라 Repository를 갖는 다른 서비스와 달리
순수 함수(모델/dataclass -> 엑셀·PDF 바이트)로만 구성한다.
"""

from datetime import datetime
from io import BytesIO
from typing import Optional

from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from models.order import Cancellation, Exchange, Order, Return
from services.report_service import MonthlyManagementReport

ORDER_EXPORT_HEADERS = ["ID", "플랫폼ID", "플랫폼주문번호", "고객ID", "상태", "주문일시", "주문금액", "할인금액"]
EXCHANGE_EXPORT_HEADERS = ["ID", "주문ID", "주문항목ID", "사유", "상태", "신청일시", "완료일시"]
RETURN_EXPORT_HEADERS = ["ID", "주문ID", "주문항목ID", "사유", "환불금액", "상태", "신청일시", "완료일시"]
CANCELLATION_EXPORT_HEADERS = ["ID", "주문ID", "사유", "환불금액", "상태", "신청일시", "완료일시"]


def _fmt(dt: Optional[datetime]) -> Optional[str]:
    return dt.strftime("%Y-%m-%d %H:%M:%S") if dt else None


def orders_to_excel(orders: list[Order]) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "주문내역"
    ws.append(ORDER_EXPORT_HEADERS)
    for order in orders:
        ws.append(
            [
                order.id,
                order.platform_id,
                order.platform_order_no,
                order.customer_id,
                order.status,
                _fmt(order.order_date),
                float(order.total_amount),
                float(order.discount_amount),
            ]
        )

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def exchanges_to_excel(exchanges: list[Exchange]) -> BytesIO:
    """SRS FR-REPORT-04: 교환내역 엑셀 내보내기."""
    wb = Workbook()
    ws = wb.active
    ws.title = "교환내역"
    ws.append(EXCHANGE_EXPORT_HEADERS)
    for e in exchanges:
        ws.append([e.id, e.order_id, e.order_item_id, e.reason, e.status, _fmt(e.requested_at), _fmt(e.completed_at)])

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def returns_to_excel(returns: list[Return]) -> BytesIO:
    """SRS FR-REPORT-04: 반품내역 엑셀 내보내기."""
    wb = Workbook()
    ws = wb.active
    ws.title = "반품내역"
    ws.append(RETURN_EXPORT_HEADERS)
    for r in returns:
        ws.append(
            [
                r.id,
                r.order_id,
                r.order_item_id,
                r.reason,
                float(r.refund_amount) if r.refund_amount is not None else None,
                r.status,
                _fmt(r.requested_at),
                _fmt(r.completed_at),
            ]
        )

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def cancellations_to_excel(cancellations: list[Cancellation]) -> BytesIO:
    """SRS FR-REPORT-04: 취소내역 엑셀 내보내기."""
    wb = Workbook()
    ws = wb.active
    ws.title = "취소내역"
    ws.append(CANCELLATION_EXPORT_HEADERS)
    for c in cancellations:
        ws.append(
            [
                c.id,
                c.order_id,
                c.reason,
                float(c.refund_amount) if c.refund_amount is not None else None,
                c.status,
                _fmt(c.requested_at),
                _fmt(c.completed_at),
            ]
        )

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def monthly_report_to_excel(report: MonthlyManagementReport) -> BytesIO:
    """SRS FR-REPORT-04: 월별손익보고서를 시트별로 나눈 엑셀로 내보낸다."""
    wb = Workbook()

    ws = wb.active
    ws.title = "손익요약"
    pl = report.profit_loss
    ws.append(["기간", report.period_key])
    ws.append(["총매출(gross)", float(pl.gross_revenue)])
    ws.append(["순매출(net)", float(pl.net_revenue)])
    ws.append(["주문건수", pl.order_count])
    ws.append(["광고비", float(pl.ad_cost)])
    ws.append(["광고전환매출", float(pl.ad_conversion_revenue)])
    ws.append(["상품원가", float(pl.cost_of_goods)])
    ws.append(["플랫폼수수료", float(pl.platform_fee)])
    ws.append(["배송비", float(pl.shipping_cost)])
    ws.append(["포장비", float(pl.packaging_cost)])
    ws.append(["기타비용", float(pl.other_cost)])
    ws.append(["총비용", float(pl.total_cost)])
    ws.append(["순이익", float(pl.net_profit)])
    ws.append(["순이익률(%)", float(pl.net_profit_rate)])
    ws.append(["반품률(%)", report.return_rate])
    ws.append(["교환률(%)", report.exchange_rate])
    ws.append(["취소율(%)", report.cancel_rate])
    ws.append(["고정비", report.cost_type_breakdown.fixed_cost])
    ws.append(["변동비", report.cost_type_breakdown.variable_cost])

    ws_platform = wb.create_sheet("플랫폼별매출")
    ws_platform.append(["플랫폼코드", "플랫폼명", "매출", "주문건수"])
    for platform_row in report.platform_breakdown:
        ws_platform.append(
            [platform_row.platform_code, platform_row.platform_name, platform_row.revenue, platform_row.order_count]
        )

    ws_best = wb.create_sheet("베스트상품")
    ws_best.append(["상품옵션ID", "판매수량", "매출", "순이익", "광고비", "ROAS(%)"])
    for product_row in report.best_products:
        ws_best.append(
            [
                product_row.product_option_id,
                product_row.sales_qty,
                float(product_row.revenue),
                float(product_row.net_profit),
                float(product_row.ad_cost),
                float(product_row.roas),
            ]
        )

    ws_worst = wb.create_sheet("워스트상품")
    ws_worst.append(["상품옵션ID", "판매수량", "매출", "순이익", "광고비", "ROAS(%)"])
    for product_row in report.worst_products:
        ws_worst.append(
            [
                product_row.product_option_id,
                product_row.sales_qty,
                float(product_row.revenue),
                float(product_row.net_profit),
                float(product_row.ad_cost),
                float(product_row.roas),
            ]
        )

    ws_ad = wb.create_sheet("광고성과보고서")
    ws_ad.append(["광고플랫폼", "노출", "클릭", "비용", "전환", "전환매출", "CPC", "CPM", "ROAS(%)"])
    for ad_row in report.ad_performance:
        ws_ad.append(
            [
                ad_row.ad_platform_code,
                ad_row.impressions,
                ad_row.clicks,
                ad_row.cost,
                ad_row.conversions,
                ad_row.conversion_amount,
                ad_row.cpc,
                ad_row.cpm,
                ad_row.roas,
            ]
        )

    if report.kpi_comparisons:
        ws_kpi = wb.create_sheet("KPI")
        ws_kpi.append(["지표", "목표", "실적", "달성률(%)"])
        for kpi in report.kpi_comparisons:
            ws_kpi.append([kpi.metric, kpi.target_value, kpi.actual_value, kpi.achievement_rate])

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def monthly_report_to_pdf(report: MonthlyManagementReport) -> BytesIO:
    """SRS FR-REPORT-04: 월별 경영보고서(손익/플랫폼별/광고성과/상품별/KPI 요약)를 PDF로 내보낸다."""
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=15 * mm, bottomMargin=15 * mm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("ReportTitle", parent=styles["Title"], fontSize=16)
    heading_style = ParagraphStyle("ReportHeading", parent=styles["Heading2"], spaceBefore=10, spaceAfter=4)

    def _table(data: list[list]) -> Table:
        table = Table(data, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#333333")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F5F5")]),
                ]
            )
        )
        return table

    elements: list[object] = [Paragraph(f"{report.period_key} 월별 경영보고서", title_style), Spacer(1, 8)]

    pl = report.profit_loss
    elements.append(Paragraph("손익 요약", heading_style))
    elements.append(
        _table(
            [
                ["항목", "값"],
                ["순매출", f"{float(pl.net_revenue):,.2f}"],
                ["주문건수", str(pl.order_count)],
                ["광고비", f"{float(pl.ad_cost):,.2f}"],
                ["상품원가", f"{float(pl.cost_of_goods):,.2f}"],
                ["플랫폼수수료", f"{float(pl.platform_fee):,.2f}"],
                ["총비용", f"{float(pl.total_cost):,.2f}"],
                ["순이익", f"{float(pl.net_profit):,.2f}"],
                ["순이익률(%)", f"{float(pl.net_profit_rate):,.2f}"],
                ["반품률(%)", f"{report.return_rate:,.2f}"],
                ["교환률(%)", f"{report.exchange_rate:,.2f}"],
                ["취소율(%)", f"{report.cancel_rate:,.2f}"],
                ["고정비", f"{report.cost_type_breakdown.fixed_cost:,.2f}"],
                ["변동비", f"{report.cost_type_breakdown.variable_cost:,.2f}"],
            ]
        )
    )

    elements.append(Paragraph("플랫폼별 매출", heading_style))
    platform_rows = [["플랫폼", "매출", "주문건수"]] + [
        [p.platform_name, f"{p.revenue:,.2f}", str(p.order_count)] for p in report.platform_breakdown
    ]
    elements.append(_table(platform_rows if len(platform_rows) > 1 else platform_rows + [["-", "-", "-"]]))

    elements.append(Paragraph("광고 성과", heading_style))
    ad_rows = [["광고플랫폼", "비용", "전환매출", "CPC", "CPM", "ROAS(%)"]] + [
        [
            a.ad_platform_code,
            f"{a.cost:,.2f}",
            f"{a.conversion_amount:,.2f}",
            f"{a.cpc:,.2f}",
            f"{a.cpm:,.2f}",
            f"{a.roas:,.2f}",
        ]
        for a in report.ad_performance
    ]
    elements.append(_table(ad_rows if len(ad_rows) > 1 else ad_rows + [["-", "-", "-", "-", "-", "-"]]))

    elements.append(Paragraph("베스트 상품 Top 5", heading_style))
    best_rows = [["상품옵션ID", "매출", "순이익", "ROAS(%)"]] + [
        [str(p.product_option_id), f"{float(p.revenue):,.2f}", f"{float(p.net_profit):,.2f}", f"{float(p.roas):,.2f}"]
        for p in report.best_products
    ]
    elements.append(_table(best_rows if len(best_rows) > 1 else best_rows + [["-", "-", "-", "-"]]))

    elements.append(Paragraph("워스트 상품 Top 5", heading_style))
    worst_rows = [["상품옵션ID", "매출", "순이익", "ROAS(%)"]] + [
        [str(p.product_option_id), f"{float(p.revenue):,.2f}", f"{float(p.net_profit):,.2f}", f"{float(p.roas):,.2f}"]
        for p in report.worst_products
    ]
    elements.append(_table(worst_rows if len(worst_rows) > 1 else worst_rows + [["-", "-", "-", "-"]]))

    if report.kpi_comparisons:
        elements.append(Paragraph("KPI 달성률", heading_style))
        kpi_rows = [["지표", "목표", "실적", "달성률(%)"]] + [
            [k.metric, f"{k.target_value:,.2f}", f"{k.actual_value:,.2f}", f"{k.achievement_rate:,.2f}"]
            for k in report.kpi_comparisons
        ]
        elements.append(_table(kpi_rows))

    doc.build(elements)
    buffer.seek(0)
    return buffer
