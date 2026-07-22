"""
tests/unit/test_export_service.py
----------------------------------------
export_service 단위 테스트. SRS FR-ORD-05(주문 엑셀 내보내기).
"""

from datetime import datetime, timezone

from openpyxl import load_workbook

from models.order import Order
from services.export_service import ORDER_EXPORT_HEADERS, orders_to_excel


def _make_order(order_id_hint: int) -> Order:
    order = Order(
        platform_id=1,
        platform_order_no=f"EXPORT-{order_id_hint}",
        customer_id=None,
        status="DELIVERED",
        order_date=datetime(2026, 1, 1, 10, 0, tzinfo=timezone.utc),
        total_amount=10000,
        discount_amount=500,
    )
    return order


class TestOrdersToExcel:
    def test_writes_header_row(self):
        buffer = orders_to_excel([])

        wb = load_workbook(buffer)
        ws = wb.active
        header = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]

        assert header == ORDER_EXPORT_HEADERS

    def test_writes_one_row_per_order(self):
        orders = [_make_order(1), _make_order(2)]

        buffer = orders_to_excel(orders)

        wb = load_workbook(buffer)
        ws = wb.active
        assert ws.max_row == 3  # 헤더 1 + 주문 2

    def test_row_values_match_order_fields(self):
        order = _make_order(1)

        buffer = orders_to_excel([order])

        wb = load_workbook(buffer)
        ws = wb.active
        row = [cell.value for cell in next(ws.iter_rows(min_row=2, max_row=2))]
        assert row[2] == "EXPORT-1"  # 플랫폼주문번호
        assert row[4] == "DELIVERED"  # 상태
        assert row[6] == 10000.0  # 주문금액
        assert row[7] == 500.0  # 할인금액
