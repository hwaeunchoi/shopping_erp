"""
tests/unit/test_report_service.py
--------------------------------------
ReportService(월별 경영보고서 조합) 및 XLSX/PDF 내보내기 단위 테스트.
SRS FR-REPORT-03/04 대응.
"""

from datetime import datetime, timezone

from openpyxl import load_workbook

from models.ad import AdCampaign, AdPerformanceDaily
from models.analytics import KpiTarget
from models.cost import Cost
from models.order import Order
from services.export_service import monthly_report_to_excel, monthly_report_to_pdf
from services.report_service import ReportService


def _make_order(db_session, platform, customer, order_date, total_amount, status="DELIVERED", order_no="RPT-ORDER"):
    order = Order(
        platform_id=platform.id,
        platform_order_no=order_no,
        customer_id=customer.id,
        status=status,
        order_date=order_date,
        total_amount=total_amount,
        discount_amount=0,
    )
    db_session.add(order)
    db_session.flush()
    return order


class TestGenerateMonthlyReport:
    def test_empty_month_returns_zeroed_report(self, db_session, platform):
        report = ReportService(db_session).generate_monthly_report(2026, 8)

        assert report.period_key == "2026-08"
        assert report.profit_loss.order_count == 0
        assert report.platform_breakdown == []
        assert report.best_products == []
        assert report.return_rate == 0.0

    def test_platform_breakdown_and_rates(self, db_session, platform, customer, product_option):
        _make_order(
            db_session, platform, customer, datetime(2026, 9, 5, 10, 0, tzinfo=timezone.utc), 10000, order_no="RPT-1"
        )
        _make_order(
            db_session,
            platform,
            customer,
            datetime(2026, 9, 6, 10, 0, tzinfo=timezone.utc),
            20000,
            status="CANCELED",
            order_no="RPT-2",
        )
        _make_order(
            db_session,
            platform,
            customer,
            datetime(2026, 9, 7, 10, 0, tzinfo=timezone.utc),
            15000,
            status="RETURNED",
            order_no="RPT-3",
        )

        report = ReportService(db_session).generate_monthly_report(2026, 9)

        assert len(report.platform_breakdown) == 1
        assert report.platform_breakdown[0].revenue == 10000
        assert report.platform_breakdown[0].order_count == 1
        # 전체 3건 중 1건 취소, 1건 반품
        assert report.cancel_rate == round(1 / 3 * 100, 2)
        assert report.return_rate == round(1 / 3 * 100, 2)

    def test_ad_performance_grouped_by_ad_platform(self, db_session, platform, customer, product_option):
        campaign = AdCampaign(
            ad_platform_code="naver_search_ad",
            platform_campaign_id="RPT-CAMPAIGN",
            product_option_id=product_option.id,
            is_active=True,
        )
        db_session.add(campaign)
        db_session.flush()
        db_session.add(
            AdPerformanceDaily(
                campaign_id=campaign.id,
                stat_date=datetime(2026, 10, 5).date(),
                impressions=1000,
                clicks=50,
                cost=10000,
                conversions=5,
                conversion_amount=50000,
            )
        )
        db_session.flush()

        report = ReportService(db_session).generate_monthly_report(2026, 10)

        assert len(report.ad_performance) == 1
        ad = report.ad_performance[0]
        assert ad.ad_platform_code == "naver_search_ad"
        assert ad.cpc == 200.0
        assert ad.roas == 500.0

    def test_kpi_comparison_computed_when_target_exists(self, db_session, platform, customer, product_option):
        _make_order(
            db_session, platform, customer, datetime(2026, 11, 1, 10, 0, tzinfo=timezone.utc), 10000, order_no="RPT-KPI"
        )
        db_session.add(
            KpiTarget(period_key="2026-11", metric="REVENUE", target_value=5000, created_at=datetime.now(timezone.utc))
        )
        db_session.flush()

        report = ReportService(db_session).generate_monthly_report(2026, 11)

        revenue_kpi = next(k for k in report.kpi_comparisons if k.metric == "REVENUE")
        assert revenue_kpi.target_value == 5000
        assert revenue_kpi.actual_value == 10000
        assert revenue_kpi.achievement_rate == 200.0

    def test_cost_type_breakdown_splits_fixed_and_variable(self, db_session, platform):
        """SRS FR-COST-02: 고정비/변동비 구분이 월별 보고서에 반영되는지 검증."""
        db_session.add(
            Cost(
                category="ETC",
                cost_type="FIXED",
                amount=50000,
                incurred_date=datetime(2026, 4, 10).date(),
                created_at=datetime.now(timezone.utc),
            )
        )
        db_session.add(
            Cost(
                category="SHIPPING",
                cost_type="VARIABLE",
                amount=3000,
                incurred_date=datetime(2026, 4, 15).date(),
                created_at=datetime.now(timezone.utc),
            )
        )
        db_session.flush()

        report = ReportService(db_session).generate_monthly_report(2026, 4)

        assert report.cost_type_breakdown.fixed_cost == 50000
        assert report.cost_type_breakdown.variable_cost == 3000


class TestExportMonthlyReport:
    def test_to_excel_produces_readable_workbook_with_expected_sheets(self, db_session, platform, customer):
        _make_order(
            db_session,
            platform,
            customer,
            datetime(2026, 12, 1, 10, 0, tzinfo=timezone.utc),
            10000,
            order_no="RPT-XLSX",
        )
        report = ReportService(db_session).generate_monthly_report(2026, 12)

        buffer = monthly_report_to_excel(report)

        wb = load_workbook(buffer)
        assert wb.sheetnames == ["손익요약", "플랫폼별매출", "베스트상품", "워스트상품", "광고성과보고서"]

    def test_to_pdf_produces_nonempty_pdf_bytes(self, db_session, platform, customer):
        _make_order(
            db_session, platform, customer, datetime(2027, 1, 1, 10, 0, tzinfo=timezone.utc), 10000, order_no="RPT-PDF"
        )
        report = ReportService(db_session).generate_monthly_report(2027, 1)

        buffer = monthly_report_to_pdf(report)

        data = buffer.getvalue()
        assert data[:4] == b"%PDF"
        assert len(data) > 500
