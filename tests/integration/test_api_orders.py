"""
tests/integration/test_api_orders.py
------------------------------------------
api/routers/orders.py 통합 테스트: 조회, 404, 인증/RBAC, 커넥터 연동 sync.
"""

from datetime import datetime, timezone


def _create_order(api_session_factory, platform_id: int, order_no: str, order_date: datetime, status: str = "NEW"):
    from models.order import Order

    db = api_session_factory()
    try:
        order = Order(
            platform_id=platform_id,
            platform_order_no=order_no,
            status=status,
            order_date=order_date,
            total_amount=10000,
            discount_amount=0,
        )
        db.add(order)
        db.commit()
        return order.id
    finally:
        db.close()


class TestListAndGet:
    def test_list_orders_empty(self, client, auth_headers):
        resp = client.get("/api/orders", headers=auth_headers)
        assert resp.status_code == 200
        body = resp.json()
        assert body["rows"] == []
        assert body["total"] == 0

    def test_get_missing_order_returns_404(self, client, auth_headers):
        resp = client.get("/api/orders/999999", headers=auth_headers)
        assert resp.status_code == 404

    def test_detail_includes_product_and_customer_info(self, client, auth_headers, api_session_factory, seed_data):
        """상품명/옵션명/SKU(판매자 내부코드)와 구매자명/연락처가 상세 조회에 포함되는지 확인한다."""
        from models.customer import Customer
        from models.order import Order, OrderItem
        from models.product import Product, ProductOption

        db = api_session_factory()
        try:
            customer = Customer(
                platform_id=seed_data["platform_id"],
                platform_customer_key="CUST-DETAIL-1",
                name="홍길동",
                phone="010-1111-2222",
            )
            db.add(customer)
            db.flush()
            product = Product(name="테스트 상품", category="테스트", base_price=10000, status="ACTIVE")
            db.add(product)
            db.flush()
            option = ProductOption(product_id=product.id, sku_code="DETAIL-SKU-1", option_name="블랙/L", is_active=True)
            db.add(option)
            db.flush()
            order = Order(
                platform_id=seed_data["platform_id"],
                platform_order_no="DETAIL-ORDER-1",
                customer_id=customer.id,
                status="NEW",
                order_date=datetime(2026, 1, 1, tzinfo=timezone.utc),
                total_amount=20000,
                discount_amount=0,
            )
            db.add(order)
            db.flush()
            db.add(
                OrderItem(
                    order_id=order.id, product_option_id=option.id, quantity=2, unit_price=10000, line_amount=20000
                )
            )
            db.commit()
            order_id = order.id
        finally:
            db.close()

        resp = client.get(f"/api/orders/{order_id}", headers=auth_headers)

        assert resp.status_code == 200
        body = resp.json()
        assert body["customer_name"] == "홍길동"
        assert body["customer_phone"] == "010-1111-2222"
        assert body["items"][0]["product_name"] == "테스트 상품"
        assert body["items"][0]["option_name"] == "블랙/L"
        assert body["items"][0]["sku_code"] == "DETAIL-SKU-1"

    def test_orders_require_authentication(self, client, seed_data):
        resp = client.get("/api/orders")
        assert resp.status_code == 401


class TestSync:
    def test_sync_unknown_platform_returns_404(self, client, auth_headers):
        resp = client.post(
            "/api/orders/sync",
            json={"platform_id": 999999, "warehouse_id": 1, "start_date": "2026-01-01", "end_date": "2026-01-02"},
            headers=auth_headers,
        )
        assert resp.status_code == 404

    def test_sync_without_credentials_returns_409(self, client, auth_headers, seed_data):
        # 인증정보가 없는 채널은 더미 주문을 성공으로 반환하지 않고 연결정보 오류(409)를 준다.
        resp = client.post(
            "/api/orders/sync",
            json={
                "platform_id": seed_data["platform_id"],
                "warehouse_id": seed_data["warehouse_id"],
                "start_date": "2026-01-01",
                "end_date": "2026-01-05",
            },
            headers=auth_headers,
        )
        assert resp.status_code == 409
        body = resp.json()
        assert body["error_code"] == "CREDENTIAL_MISSING"
        assert body["retryable"] is False
        assert "secret" not in str(body).lower()

    def test_sync_unverified_channel_returns_501(self, client, auth_headers, api_session_factory):
        # 미검증 채널(ESM 등)은 팩토리에서 미지원 오류 -> 501(더미 유입 없음).
        from models.platform import Platform

        db = api_session_factory()
        try:
            esm = Platform(code="esm_x", name="ESM", connector_class="EsmConnector", is_active=True)
            db.add(esm)
            db.commit()
            esm_id = esm.id
        finally:
            db.close()

        resp = client.post(
            "/api/orders/sync",
            json={"platform_id": esm_id, "warehouse_id": 1, "start_date": "2026-01-01", "end_date": "2026-01-02"},
            headers=auth_headers,
        )
        assert resp.status_code == 501
        assert resp.json()["error_code"] == "CAPABILITY_UNSUPPORTED"

    def test_sync_external_api_retryable_returns_503(self, client, auth_headers, seed_data, monkeypatch):
        from integrations.malls.errors import MarketplaceExternalAPIError

        class _BoomConnector:
            def fetch_orders(self, *a, **k):
                raise MarketplaceExternalAPIError("coupang", "SERVER_ERROR", True, http_status=500)

        monkeypatch.setattr("api.routers.orders.get_mall_connector", lambda *a, **k: _BoomConnector())
        resp = client.post(
            "/api/orders/sync",
            json={
                "platform_id": seed_data["platform_id"],
                "warehouse_id": seed_data["warehouse_id"],
                "start_date": "2026-01-01",
                "end_date": "2026-01-02",
            },
            headers=auth_headers,
        )
        assert resp.status_code == 503
        body = resp.json()
        assert body["error_code"] == "EXTERNAL_API_ERROR"
        assert body["retryable"] is True

    def test_sync_external_api_non_retryable_returns_502(self, client, auth_headers, seed_data, monkeypatch):
        from integrations.malls.errors import MarketplaceExternalAPIError

        class _BoomConnector:
            def fetch_orders(self, *a, **k):
                raise MarketplaceExternalAPIError("coupang", "AUTH_FAILED", False, http_status=401)

        monkeypatch.setattr("api.routers.orders.get_mall_connector", lambda *a, **k: _BoomConnector())
        resp = client.post(
            "/api/orders/sync",
            json={
                "platform_id": seed_data["platform_id"],
                "warehouse_id": seed_data["warehouse_id"],
                "start_date": "2026-01-01",
                "end_date": "2026-01-02",
            },
            headers=auth_headers,
        )
        assert resp.status_code == 502
        assert resp.json()["retryable"] is False

    def test_sync_requires_order_edit_permission(self, client, api_session_factory, seed_data):
        db = api_session_factory()
        try:
            from core.security import hash_password
            from models.user import Permission, Role, RolePermission, User

            role = Role(name="Viewer")
            db.add(role)
            db.flush()
            view_perm = db.query(Permission).filter_by(code="ORDER_VIEW").first()
            db.add(RolePermission(role_id=role.id, permission_id=view_perm.id))
            db.add(
                User(
                    username="viewer",
                    password_hash=hash_password("pw123456"),
                    name="뷰어",
                    role_id=role.id,
                    is_active=True,
                )
            )
            db.commit()
        finally:
            db.close()

        login = client.post("/api/auth/login", data={"username": "viewer", "password": "pw123456"})
        headers = {"Authorization": f"Bearer {login.json()['access_token']}"}

        allowed = client.get("/api/orders", headers=headers)
        assert allowed.status_code == 200

        forbidden = client.post(
            "/api/orders/sync",
            json={
                "platform_id": seed_data["platform_id"],
                "warehouse_id": seed_data["warehouse_id"],
                "start_date": "2026-01-01",
                "end_date": "2026-01-02",
            },
            headers=headers,
        )
        assert forbidden.status_code == 403


class _StubClaimConnector:
    """API 테스트용 클레임 스텁(capability·기능별 데이터/예외 제어, 실 네트워크 없음)."""

    def __init__(self, *, supports=(), errors=None, data=None):
        self.supports_cancellation_sync = "cancellation" in supports
        self.supports_return_sync = "return" in supports
        self.supports_exchange_sync = "exchange" in supports
        self._errors = errors or {}
        self._data = data or {}

    def fetch_cancellations(self, s, e):
        if "cancellations" in self._errors:
            raise self._errors["cancellations"]
        return self._data.get("cancellations", [])

    def fetch_returns(self, s, e):
        if "returns" in self._errors:
            raise self._errors["returns"]
        return self._data.get("returns", [])

    def fetch_exchanges(self, s, e):
        if "exchanges" in self._errors:
            raise self._errors["exchanges"]
        return self._data.get("exchanges", [])


class TestSyncClaims:
    def _post(self, client, auth_headers, seed_data):
        return client.post(
            "/api/orders/sync-claims",
            json={
                "platform_id": seed_data["platform_id"],
                "warehouse_id": seed_data["warehouse_id"],
                "start_date": "2026-01-01",
                "end_date": "2026-01-31",
            },
            headers=auth_headers,
        )

    def _patch(self, monkeypatch, conn):
        monkeypatch.setattr("api.routers.orders.get_mall_connector", lambda *a, **k: conn)

    def test_all_unsupported_returns_501_with_structured_body(self, client, auth_headers, seed_data):
        # seed 플랫폼(쿠팡)은 세 capability 모두 False -> 전 기능 미지원 -> 501(성공 위장 아님).
        resp = self._post(client, auth_headers, seed_data)
        assert resp.status_code == 501
        body = resp.json()
        assert body["overall_status"] == "UNSUPPORTED"
        assert body["cancellations"]["status"] == "UNSUPPORTED"
        assert body["returns"]["status"] == "UNSUPPORTED"
        assert body["exchanges"]["status"] == "UNSUPPORTED"
        assert "secret" not in str(body).lower()

    def test_success_returns_200(self, client, auth_headers, seed_data, monkeypatch):
        self._patch(monkeypatch, _StubClaimConnector(supports=("cancellation", "return", "exchange")))
        resp = self._post(client, auth_headers, seed_data)
        assert resp.status_code == 200
        body = resp.json()
        assert body["overall_status"] == "SUCCESS"
        # 지원+0건은 SUCCESS(count 0) - 미지원과 구분됨
        assert body["cancellations"] == {"status": "SUCCESS", "count": 0, "reason_code": None, "retryable": None}

    def test_partial_returns_200(self, client, auth_headers, seed_data, monkeypatch):
        from integrations.malls.errors import MarketplaceExternalAPIError

        conn = _StubClaimConnector(
            supports=("cancellation", "return"),
            errors={"returns": MarketplaceExternalAPIError("x", "SERVER_ERROR", True)},
        )
        self._patch(monkeypatch, conn)
        resp = self._post(client, auth_headers, seed_data)
        assert resp.status_code == 200
        assert resp.json()["overall_status"] == "PARTIAL"

    def test_all_credential_missing_returns_409(self, client, auth_headers, seed_data, monkeypatch):
        from integrations.malls.errors import MarketplaceCredentialMissingError

        conn = _StubClaimConnector(
            supports=("cancellation", "return", "exchange"),
            errors={k: MarketplaceCredentialMissingError("x") for k in ("cancellations", "returns", "exchanges")},
        )
        self._patch(monkeypatch, conn)
        resp = self._post(client, auth_headers, seed_data)
        assert resp.status_code == 409
        assert resp.json()["overall_status"] == "FAILED"

    def test_all_external_retryable_returns_503(self, client, auth_headers, seed_data, monkeypatch):
        from integrations.malls.errors import MarketplaceExternalAPIError

        conn = _StubClaimConnector(
            supports=("cancellation", "return", "exchange"),
            errors={
                k: MarketplaceExternalAPIError("x", "SERVER_ERROR", True)
                for k in ("cancellations", "returns", "exchanges")
            },
        )
        self._patch(monkeypatch, conn)
        resp = self._post(client, auth_headers, seed_data)
        assert resp.status_code == 503
        # 비200에서도 구조화 body 유지
        assert resp.json()["cancellations"]["reason_code"] == "SERVER_ERROR"

    def test_all_external_non_retryable_returns_502(self, client, auth_headers, seed_data, monkeypatch):
        from integrations.malls.errors import MarketplaceExternalAPIError

        conn = _StubClaimConnector(
            supports=("cancellation", "return", "exchange"),
            errors={
                k: MarketplaceExternalAPIError("x", "AUTH_FAILED", False)
                for k in ("cancellations", "returns", "exchanges")
            },
        )
        self._patch(monkeypatch, conn)
        resp = self._post(client, auth_headers, seed_data)
        assert resp.status_code == 502
        assert resp.json()["overall_status"] == "FAILED"


class TestFilters:
    def test_filters_by_platform_id(self, client, auth_headers, api_session_factory, seed_data):
        from models.platform import Platform

        db = api_session_factory()
        try:
            other_platform = Platform(code="other_platform", name="다른플랫폼", connector_class="X", is_active=True)
            db.add(other_platform)
            db.commit()
            other_platform_id = other_platform.id
        finally:
            db.close()

        _create_order(api_session_factory, seed_data["platform_id"], "FLT-1", datetime(2026, 5, 1, tzinfo=timezone.utc))
        _create_order(api_session_factory, other_platform_id, "FLT-2", datetime(2026, 5, 1, tzinfo=timezone.utc))

        resp = client.get(f"/api/orders?platform_id={other_platform_id}", headers=auth_headers)

        assert resp.status_code == 200
        body = resp.json()["rows"]
        assert len(body) == 1
        assert body[0]["platform_order_no"] == "FLT-2"

    def test_filters_by_date_range(self, client, auth_headers, api_session_factory, seed_data):
        _create_order(
            api_session_factory, seed_data["platform_id"], "FLT-IN", datetime(2026, 6, 15, tzinfo=timezone.utc)
        )
        _create_order(
            api_session_factory, seed_data["platform_id"], "FLT-OUT", datetime(2026, 7, 15, tzinfo=timezone.utc)
        )

        resp = client.get("/api/orders?start_date=2026-06-01&end_date=2026-06-30", headers=auth_headers)

        assert resp.status_code == 200
        order_nos = [o["platform_order_no"] for o in resp.json()["rows"]]
        assert "FLT-IN" in order_nos
        assert "FLT-OUT" not in order_nos

    def test_combines_status_and_platform_filters(self, client, auth_headers, api_session_factory, seed_data):
        _create_order(
            api_session_factory,
            seed_data["platform_id"],
            "FLT-STATUS-1",
            datetime(2026, 5, 2, tzinfo=timezone.utc),
            status="DELIVERED",
        )
        _create_order(
            api_session_factory,
            seed_data["platform_id"],
            "FLT-STATUS-2",
            datetime(2026, 5, 2, tzinfo=timezone.utc),
            status="NEW",
        )

        resp = client.get(
            f"/api/orders?status_filter=DELIVERED&platform_id={seed_data['platform_id']}", headers=auth_headers
        )

        assert resp.status_code == 200
        order_nos = [o["platform_order_no"] for o in resp.json()["rows"]]
        assert "FLT-STATUS-1" in order_nos
        assert "FLT-STATUS-2" not in order_nos

    def test_keyword_searches_order_number(self, client, auth_headers, api_session_factory, seed_data):
        _create_order(
            api_session_factory, seed_data["platform_id"], "KEYWORD-UNIQUE-1", datetime(2026, 5, 3, tzinfo=timezone.utc)
        )
        _create_order(
            api_session_factory, seed_data["platform_id"], "OTHER-2", datetime(2026, 5, 3, tzinfo=timezone.utc)
        )

        resp = client.get("/api/orders?keyword=UNIQUE", headers=auth_headers)

        assert resp.status_code == 200
        order_nos = [o["platform_order_no"] for o in resp.json()["rows"]]
        assert "KEYWORD-UNIQUE-1" in order_nos
        assert "OTHER-2" not in order_nos

    def test_keyword_searches_customer_name(self, client, auth_headers, api_session_factory, seed_data):
        from models.customer import Customer
        from models.order import Order

        db = api_session_factory()
        try:
            customer = Customer(platform_id=seed_data["platform_id"], platform_customer_key="CUST-KW", name="김철수")
            db.add(customer)
            db.flush()
            db.add(
                Order(
                    platform_id=seed_data["platform_id"],
                    platform_order_no="CUST-KW-ORDER",
                    customer_id=customer.id,
                    status="NEW",
                    order_date=datetime(2026, 5, 3, tzinfo=timezone.utc),
                    total_amount=10000,
                    discount_amount=0,
                )
            )
            db.commit()
        finally:
            db.close()

        resp = client.get("/api/orders?keyword=%EA%B9%80%EC%B2%A0%EC%88%98", headers=auth_headers)

        assert resp.status_code == 200
        body = resp.json()["rows"]
        assert any(o["platform_order_no"] == "CUST-KW-ORDER" for o in body)
        matched = next(o for o in body if o["platform_order_no"] == "CUST-KW-ORDER")
        assert matched["customer_name"] == "김철수"

    def test_list_includes_platform_name(self, client, auth_headers, api_session_factory, seed_data):
        _create_order(
            api_session_factory, seed_data["platform_id"], "PLAT-NAME-1", datetime(2026, 5, 4, tzinfo=timezone.utc)
        )

        resp = client.get(f"/api/orders?platform_id={seed_data['platform_id']}", headers=auth_headers)

        assert resp.status_code == 200
        matched = next(o for o in resp.json()["rows"] if o["platform_order_no"] == "PLAT-NAME-1")
        assert matched["platform_name"] == "쿠팡"


def _create_order_with_item(api_session_factory, platform_id, order_no, order_date, product_name, sku, status="NEW"):
    from models.order import Order, OrderItem
    from models.product import Product, ProductOption

    db = api_session_factory()
    try:
        product = Product(name=product_name, category="테스트", base_price=10000, status="ACTIVE")
        db.add(product)
        db.flush()
        option = ProductOption(product_id=product.id, sku_code=sku, option_name="기본", is_active=True)
        db.add(option)
        db.flush()
        order = Order(
            platform_id=platform_id,
            platform_order_no=order_no,
            status=status,
            order_date=order_date,
            total_amount=10000,
            discount_amount=0,
        )
        db.add(order)
        db.flush()
        db.add(
            OrderItem(order_id=order.id, product_option_id=option.id, quantity=2, unit_price=5000, line_amount=10000)
        )
        db.commit()
        return order.id, option.id
    finally:
        db.close()


class TestOrderView:
    def test_row_includes_derived_statuses_and_product(self, client, auth_headers, api_session_factory, seed_data):
        order_id, _ = _create_order_with_item(
            api_session_factory,
            seed_data["platform_id"],
            "VIEW-1",
            datetime(2026, 5, 5, tzinfo=timezone.utc),
            "뷰테스트 상품",
            "VIEW-SKU-1",
        )
        resp = client.get("/api/orders?keyword=VIEW-1", headers=auth_headers)
        assert resp.status_code == 200
        row = next(r for r in resp.json()["rows"] if r["id"] == order_id)
        assert row["payment_status"] == "UNPAID"
        assert row["shipping_status"] == "UNSHIPPED"
        assert row["cs_status"] == "NONE"
        assert row["product_name"] == "뷰테스트 상품"
        assert row["sku_code"] == "VIEW-SKU-1"
        assert row["total_quantity"] == 2

    def test_keyword_searches_product_name(self, client, auth_headers, api_session_factory, seed_data):
        _create_order_with_item(
            api_session_factory,
            seed_data["platform_id"],
            "PRODKW-1",
            datetime(2026, 5, 5, tzinfo=timezone.utc),
            "유니크상품명XYZ",
            "PRODKW-SKU",
        )
        resp = client.get("/api/orders?keyword=XYZ", headers=auth_headers)
        assert resp.status_code == 200
        assert any(r["platform_order_no"] == "PRODKW-1" for r in resp.json()["rows"])

    def test_keyword_searches_sku(self, client, auth_headers, api_session_factory, seed_data):
        _create_order_with_item(
            api_session_factory,
            seed_data["platform_id"],
            "SKUKW-1",
            datetime(2026, 5, 5, tzinfo=timezone.utc),
            "SKU검색상품",
            "FINDME-SKU-99",
        )
        resp = client.get("/api/orders?keyword=FINDME", headers=auth_headers)
        assert resp.status_code == 200
        assert any(r["platform_order_no"] == "SKUKW-1" for r in resp.json()["rows"])

    def test_filter_by_sku(self, client, auth_headers, api_session_factory, seed_data):
        _create_order_with_item(
            api_session_factory,
            seed_data["platform_id"],
            "SKUF-1",
            datetime(2026, 5, 6, tzinfo=timezone.utc),
            "SKU필터상품",
            "FILTER-SKU-A",
        )
        _create_order_with_item(
            api_session_factory,
            seed_data["platform_id"],
            "SKUF-2",
            datetime(2026, 5, 6, tzinfo=timezone.utc),
            "다른상품",
            "OTHER-SKU-B",
        )
        resp = client.get("/api/orders?sku=FILTER-SKU-A", headers=auth_headers)
        assert resp.status_code == 200
        nos = [r["platform_order_no"] for r in resp.json()["rows"]]
        assert "SKUF-1" in nos
        assert "SKUF-2" not in nos

    def test_pagination(self, client, auth_headers, api_session_factory, seed_data):
        for i in range(5):
            _create_order(
                api_session_factory, seed_data["platform_id"], f"PAGE-{i}", datetime(2026, 5, 7, tzinfo=timezone.utc)
            )
        resp = client.get("/api/orders?keyword=PAGE-&page=1&page_size=2", headers=auth_headers)
        assert resp.status_code == 200
        body = resp.json()
        assert body["total"] == 5
        assert len(body["rows"]) == 2
        assert body["page_size"] == 2

    def test_detail_returns_full_bundle(self, client, auth_headers, api_session_factory, seed_data):
        order_id, _ = _create_order_with_item(
            api_session_factory,
            seed_data["platform_id"],
            "DETAIL-1",
            datetime(2026, 5, 8, tzinfo=timezone.utc),
            "상세테스트 상품",
            "DETAIL-SKU-1",
        )
        resp = client.get(f"/api/orders/{order_id}/detail", headers=auth_headers)
        assert resp.status_code == 200
        body = resp.json()
        assert body["id"] == order_id
        assert len(body["items"]) == 1
        assert body["items"][0]["product_name"] == "상세테스트 상품"
        assert "inventory" in body
        assert "status_history" in body
        assert body["shipment"] is None

    def test_detail_missing_returns_404(self, client, auth_headers):
        resp = client.get("/api/orders/999999/detail", headers=auth_headers)
        assert resp.status_code == 404


class TestBulkActions:
    def test_bulk_ship_registers_tracking(self, client, auth_headers, api_session_factory, seed_data):
        order_id = _create_order(
            api_session_factory, seed_data["platform_id"], "BULK-SHIP-1", datetime(2026, 5, 9, tzinfo=timezone.utc)
        )
        resp = client.post(
            "/api/orders/bulk/ship",
            json={"items": [{"order_id": order_id, "carrier": "CJ대한통운", "tracking_no": "1234567890"}]},
            headers=auth_headers,
        )
        assert resp.status_code == 200
        assert resp.json()["succeeded"] == 1
        detail = client.get(f"/api/orders/{order_id}/detail", headers=auth_headers).json()
        assert detail["shipment"]["tracking_no"] == "1234567890"
        assert detail["shipping_status"] == "SHIPPING"

    def test_bulk_status_changes_and_logs(self, client, auth_headers, api_session_factory, seed_data):
        order_id = _create_order(
            api_session_factory, seed_data["platform_id"], "BULK-ST-1", datetime(2026, 5, 9, tzinfo=timezone.utc)
        )
        resp = client.post(
            "/api/orders/bulk/status", json={"order_ids": [order_id], "status": "PREPARING"}, headers=auth_headers
        )
        assert resp.status_code == 200
        assert resp.json()["succeeded"] == 1
        detail = client.get(f"/api/orders/{order_id}/detail", headers=auth_headers).json()
        assert detail["order_status"] == "PREPARING"
        assert any(h["to_status"] == "PREPARING" for h in detail["status_history"])

    def test_bulk_memo_adds_to_all(self, client, auth_headers, api_session_factory, seed_data):
        id1 = _create_order(
            api_session_factory, seed_data["platform_id"], "BULK-M-1", datetime(2026, 5, 9, tzinfo=timezone.utc)
        )
        id2 = _create_order(
            api_session_factory, seed_data["platform_id"], "BULK-M-2", datetime(2026, 5, 9, tzinfo=timezone.utc)
        )
        resp = client.post(
            "/api/orders/bulk/memo", json={"order_ids": [id1, id2], "content": "일괄 메모 테스트"}, headers=auth_headers
        )
        assert resp.status_code == 200
        assert resp.json()["succeeded"] == 2
        d1 = client.get(f"/api/orders/{id1}/detail", headers=auth_headers).json()
        assert any(m["content"] == "일괄 메모 테스트" for m in d1["memos"])

    def test_bulk_ship_reports_missing_order(self, client, auth_headers):
        resp = client.post(
            "/api/orders/bulk/ship",
            json={"items": [{"order_id": 999999, "carrier": "X", "tracking_no": "Y"}]},
            headers=auth_headers,
        )
        assert resp.status_code == 200
        body = resp.json()
        assert body["succeeded"] == 0
        assert body["failed"] == 1


class TestWorkbenchHub:
    def test_bulk_assign_and_filter_by_assignee(self, client, auth_headers, api_session_factory, seed_data):
        oid = _create_order(
            api_session_factory, seed_data["platform_id"], "ASSIGN-1", datetime(2026, 5, 10, tzinfo=timezone.utc)
        )
        admin_id = seed_data["admin_id"]
        resp = client.patch(
            "/api/orders/bulk/assign", json={"order_ids": [oid], "assignee_id": admin_id}, headers=auth_headers
        )
        assert resp.status_code == 200
        assert resp.json()["succeeded"] == 1

        listed = client.get(f"/api/orders?assignee_id={admin_id}", headers=auth_headers)
        row = next(r for r in listed.json()["rows"] if r["id"] == oid)
        assert row["assignee_id"] == admin_id
        assert row["assignee_name"] == "관리자"

    def test_create_cs_from_workbench_and_status_reflects(self, client, auth_headers, api_session_factory, seed_data):
        oid = _create_order(
            api_session_factory, seed_data["platform_id"], "CS-1", datetime(2026, 5, 10, tzinfo=timezone.utc)
        )
        resp = client.post(
            f"/api/orders/{oid}/cs", json={"type": "RETURN", "reason": "상품 불량"}, headers=auth_headers
        )
        assert resp.status_code == 201
        assert resp.json()["type"] == "RETURN"

        detail = client.get(f"/api/orders/{oid}/detail", headers=auth_headers).json()
        assert detail["cs_status"] == "RETURN"
        assert len(detail["cs_history"]) == 1
        assert detail["cs_history"][0]["type"] == "RETURN"

    def test_filter_by_cs_status(self, client, auth_headers, api_session_factory, seed_data):
        cs_order = _create_order(
            api_session_factory, seed_data["platform_id"], "CS-FLT-1", datetime(2026, 5, 10, tzinfo=timezone.utc)
        )
        _create_order(
            api_session_factory, seed_data["platform_id"], "CS-FLT-2", datetime(2026, 5, 10, tzinfo=timezone.utc)
        )
        client.post(f"/api/orders/{cs_order}/cs", json={"type": "CANCEL", "reason": "변심"}, headers=auth_headers)

        resp = client.get("/api/orders?cs_status=CANCEL", headers=auth_headers)
        nos = [r["platform_order_no"] for r in resp.json()["rows"]]
        assert "CS-FLT-1" in nos
        assert "CS-FLT-2" not in nos

    def test_create_cs_invalid_type_returns_400(self, client, auth_headers, api_session_factory, seed_data):
        oid = _create_order(
            api_session_factory, seed_data["platform_id"], "CS-BAD-1", datetime(2026, 5, 10, tzinfo=timezone.utc)
        )
        resp = client.post(f"/api/orders/{oid}/cs", json={"type": "REFUND"}, headers=auth_headers)
        assert resp.status_code == 400

    def test_detail_includes_contribution_margin(self, client, auth_headers, api_session_factory, seed_data):
        order_id, _ = _create_order_with_item(
            api_session_factory,
            seed_data["platform_id"],
            "CM-1",
            datetime(2026, 5, 11, tzinfo=timezone.utc),
            "공헌이익 상품",
            "CM-SKU-1",
        )
        detail = client.get(f"/api/orders/{order_id}/detail", headers=auth_headers).json()
        assert "contribution" in detail
        assert "contribution_margin" in detail["contribution"]
        assert "purchase_links" in detail
        assert "settlement" in detail


class TestRates:
    def test_returns_zero_rates_without_data(self, client, auth_headers, seed_data):
        resp = client.get("/api/orders/rates?start_date=2026-08-01&end_date=2026-08-31", headers=auth_headers)

        assert resp.status_code == 200
        body = resp.json()
        assert body["order_count"] == 0
        assert body["exchange_rate"] == 0.0

    def test_calculates_rates_with_data(self, client, auth_headers, api_session_factory, seed_data):
        # 취소 신청은 POST 시점(now)에 requested_at이 찍히므로, 주문일도 오늘로 맞춰서
        # order_count/cancellation_count가 같은 기간(오늘) 안에서 계산되게 한다.
        today = datetime.now(timezone.utc)
        order_id = _create_order(api_session_factory, seed_data["platform_id"], "RATE-API-1", today)
        _create_order(api_session_factory, seed_data["platform_id"], "RATE-API-2", today)
        cancel_resp = client.post(
            "/api/cancellations", json={"order_id": order_id, "reason": "단순 변심"}, headers=auth_headers
        )
        assert cancel_resp.status_code == 201

        today_str = today.date().isoformat()
        resp = client.get(f"/api/orders/rates?start_date={today_str}&end_date={today_str}", headers=auth_headers)

        assert resp.status_code == 200
        body = resp.json()
        assert body["order_count"] == 2
        assert body["cancellation_count"] == 1
        assert body["cancellation_rate"] == 50.0

    def test_requires_authentication(self, client, seed_data):
        resp = client.get("/api/orders/rates")
        assert resp.status_code == 401


class TestAlerts:
    def test_returns_zero_alerts_without_data(self, client, auth_headers, seed_data):
        resp = client.get("/api/orders/alerts", headers=auth_headers)

        assert resp.status_code == 200
        body = resp.json()
        assert body == {
            "unshipped_count": 0,
            "delayed_unshipped_count": 0,
            "exchange_pending_count": 0,
            "return_pending_count": 0,
            "cancellation_pending_count": 0,
        }

    def test_counts_unshipped_orders(self, client, auth_headers, api_session_factory, seed_data):
        _create_order(
            api_session_factory, seed_data["platform_id"], "ALERT-API-1", datetime.now(timezone.utc), status="NEW"
        )
        _create_order(
            api_session_factory, seed_data["platform_id"], "ALERT-API-2", datetime.now(timezone.utc), status="DELIVERED"
        )

        resp = client.get("/api/orders/alerts", headers=auth_headers)

        assert resp.status_code == 200
        assert resp.json()["unshipped_count"] == 1

    def test_requires_authentication(self, client, seed_data):
        resp = client.get("/api/orders/alerts")
        assert resp.status_code == 401


class TestExport:
    def test_export_returns_xlsx_with_correct_content(self, client, auth_headers, api_session_factory, seed_data):
        from io import BytesIO

        from openpyxl import load_workbook

        _create_order(
            api_session_factory, seed_data["platform_id"], "EXPORT-API-1", datetime(2026, 10, 1, tzinfo=timezone.utc)
        )

        resp = client.get("/api/orders/export?start_date=2026-10-01&end_date=2026-10-01", headers=auth_headers)

        assert resp.status_code == 200
        assert resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        assert "attachment" in resp.headers["content-disposition"]

        wb = load_workbook(BytesIO(resp.content))
        ws = wb.active
        assert ws.max_row == 2  # 헤더 1 + 주문 1
        assert ws.cell(row=2, column=3).value == "EXPORT-API-1"

    def test_requires_authentication(self, client, seed_data):
        resp = client.get("/api/orders/export")
        assert resp.status_code == 401


class TestMemos:
    def test_create_and_list_memos(self, client, auth_headers, api_session_factory, seed_data):
        order_id = _create_order(
            api_session_factory, seed_data["platform_id"], "MEMO-ORDER-1", datetime(2026, 1, 1, tzinfo=timezone.utc)
        )

        create = client.post(
            f"/api/orders/{order_id}/memos", json={"content": "고객 요청: 부재시 문앞에 놔주세요"}, headers=auth_headers
        )

        assert create.status_code == 201
        assert create.json()["content"] == "고객 요청: 부재시 문앞에 놔주세요"
        assert create.json()["target_type"] == "ORDER"
        assert create.json()["target_id"] == order_id

        listing = client.get(f"/api/orders/{order_id}/memos", headers=auth_headers)
        assert listing.status_code == 200
        assert len(listing.json()) == 1

    def test_create_memo_for_missing_order_returns_404(self, client, auth_headers):
        resp = client.post("/api/orders/999999/memos", json={"content": "메모"}, headers=auth_headers)
        assert resp.status_code == 404

    def test_list_memos_for_missing_order_returns_404(self, client, auth_headers):
        resp = client.get("/api/orders/999999/memos", headers=auth_headers)
        assert resp.status_code == 404

    def test_requires_authentication(self, client, seed_data):
        resp = client.get("/api/orders/1/memos")
        assert resp.status_code == 401

    def test_create_requires_order_edit_permission(self, client, api_session_factory, seed_data):
        from core.security import hash_password
        from models.user import Permission, Role, RolePermission, User

        order_id = _create_order(
            api_session_factory, seed_data["platform_id"], "MEMO-ORDER-2", datetime(2026, 1, 1, tzinfo=timezone.utc)
        )

        db = api_session_factory()
        try:
            role = Role(name="ViewerMemo")
            db.add(role)
            db.flush()
            view_perm = db.query(Permission).filter_by(code="ORDER_VIEW").first()
            db.add(RolePermission(role_id=role.id, permission_id=view_perm.id))
            db.add(
                User(
                    username="viewer_memo",
                    password_hash=hash_password("pw123456"),
                    name="뷰어",
                    role_id=role.id,
                    is_active=True,
                )
            )
            db.commit()
        finally:
            db.close()

        login = client.post("/api/auth/login", data={"username": "viewer_memo", "password": "pw123456"})
        headers = {"Authorization": f"Bearer {login.json()['access_token']}"}

        resp = client.post(f"/api/orders/{order_id}/memos", json={"content": "메모"}, headers=headers)
        assert resp.status_code == 403
