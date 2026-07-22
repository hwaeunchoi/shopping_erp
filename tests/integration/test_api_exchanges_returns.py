"""
tests/integration/test_api_exchanges_returns.py
------------------------------------------------------
api/routers/exchanges.py, returns.py, cancellations.py 통합 테스트.
세 라우터가 EXCHANGE_RETURN_MANAGE 권한 하나로 보호되고, 필터/페이지네이션
모양이 동일하므로 한 파일에서 함께 다룬다.
"""

from datetime import datetime, timezone

from models.order import Order, OrderItem


def _create_order_with_item(api_session_factory, seed_data, order_no: str, product_option_id: int) -> dict:
    db = api_session_factory()
    try:
        order = Order(
            platform_id=seed_data["platform_id"],
            platform_order_no=order_no,
            status="DELIVERED",
            order_date=datetime.now(timezone.utc),
            total_amount=10000,
            discount_amount=0,
        )
        db.add(order)
        db.flush()
        item = OrderItem(
            order_id=order.id,
            product_option_id=product_option_id,
            quantity=2,
            unit_price=5000,
            cost_price_snapshot=2000,
            line_amount=10000,
        )
        db.add(item)
        db.commit()
        return {"order_id": order.id, "item_id": item.id}
    finally:
        db.close()


def _seed_product_option(api_session_factory, seed_data=None, with_inventory: bool = True) -> int:
    from models.inventory import Inventory
    from models.product import Product, ProductOption

    db = api_session_factory()
    try:
        product = Product(name="테스트상품", category="test", base_price=1000, status="ACTIVE")
        db.add(product)
        db.flush()
        option = ProductOption(product_id=product.id, sku_code=f"API-TEST-SKU-{product.id}", is_active=True)
        db.add(option)
        db.flush()
        if with_inventory and seed_data is not None:
            db.add(
                Inventory(
                    product_option_id=option.id,
                    warehouse_id=seed_data["warehouse_id"],
                    sellable_stock=100,
                    reserved_stock=0,
                    safety_stock=5,
                    updated_at=datetime.now(timezone.utc),
                )
            )
        db.commit()
        return option.id
    finally:
        db.close()


class TestExchanges:
    def test_full_flow(self, client, auth_headers, api_session_factory, seed_data):
        option_id = _seed_product_option(api_session_factory)
        ctx = _create_order_with_item(api_session_factory, seed_data, "EX-API-1", option_id)

        created = client.post(
            "/api/exchanges",
            json={"order_id": ctx["order_id"], "order_item_id": ctx["item_id"], "reason": "사이즈 교환"},
            headers=auth_headers,
        )
        assert created.status_code == 201
        exchange_id = created.json()["id"]
        assert created.json()["status"] == "REQUESTED"

        got = client.get(f"/api/exchanges/{exchange_id}", headers=auth_headers)
        assert got.status_code == 200

        changed = client.patch(
            f"/api/exchanges/{exchange_id}/status", json={"status": "COMPLETED"}, headers=auth_headers
        )
        assert changed.status_code == 200
        assert changed.json()["status"] == "COMPLETED"

        order_resp = client.get(f"/api/orders/{ctx['order_id']}", headers=auth_headers)
        assert order_resp.json()["status"] == "EXCHANGED"

    def test_create_unknown_order_returns_404(self, client, auth_headers):
        resp = client.post("/api/exchanges", json={"order_id": 999999, "reason": "x"}, headers=auth_headers)
        assert resp.status_code == 404

    def test_requires_authentication(self, client, seed_data):
        resp = client.get("/api/exchanges")
        assert resp.status_code == 401


class TestReturns:
    def test_full_flow_requested_to_refunded(self, client, auth_headers, api_session_factory, seed_data):
        option_id = _seed_product_option(api_session_factory, seed_data)
        ctx = _create_order_with_item(api_session_factory, seed_data, "RET-API-1", option_id)

        created = client.post(
            "/api/returns",
            json={
                "order_id": ctx["order_id"],
                "order_item_id": ctx["item_id"],
                "reason": "상품 불량",
                "refund_amount": 10000,
            },
            headers=auth_headers,
        )
        assert created.status_code == 201
        return_id = created.json()["id"]

        received = client.patch(
            f"/api/returns/{return_id}/status",
            json={"status": "RECEIVED", "warehouse_id": seed_data["warehouse_id"]},
            headers=auth_headers,
        )
        assert received.status_code == 200
        assert received.json()["status"] == "RECEIVED"

        refunded = client.patch(
            f"/api/returns/{return_id}/status",
            json={"status": "REFUNDED", "warehouse_id": seed_data["warehouse_id"]},
            headers=auth_headers,
        )
        assert refunded.status_code == 200

        order_resp = client.get(f"/api/orders/{ctx['order_id']}", headers=auth_headers)
        assert order_resp.json()["status"] == "REFUNDED"

    def test_list_filters_by_status(self, client, auth_headers, api_session_factory, seed_data):
        option_id = _seed_product_option(api_session_factory, seed_data)
        ctx = _create_order_with_item(api_session_factory, seed_data, "RET-API-2", option_id)
        client.post(
            "/api/returns",
            json={"order_id": ctx["order_id"], "order_item_id": ctx["item_id"], "reason": "단순 변심"},
            headers=auth_headers,
        )

        resp = client.get("/api/returns?status_filter=REQUESTED", headers=auth_headers)

        assert resp.status_code == 200
        assert resp.json()["total"] >= 1
        assert all(r["status"] == "REQUESTED" for r in resp.json()["items"])

    def test_received_without_inventory_record_succeeds(self, client, auth_headers, api_session_factory, seed_data):
        """정책 5: RECEIVED는 재고를 건드리지 않으므로 재고 레코드가 없어도 성공한다.

        이전에는 RECEIVED가 즉시 재고를 복원해 재고 레코드가 없으면 404였다.
        재고 반영은 검수 시점(InventoryService.inspect_return())으로 옮겨졌다.
        """
        option_id = _seed_product_option(api_session_factory, seed_data, with_inventory=False)
        ctx = _create_order_with_item(api_session_factory, seed_data, "RET-API-3", option_id)
        created = client.post(
            "/api/returns",
            json={"order_id": ctx["order_id"], "order_item_id": ctx["item_id"], "reason": "상품 불량"},
            headers=auth_headers,
        )

        resp = client.patch(
            f"/api/returns/{created.json()['id']}/status",
            json={"status": "RECEIVED", "warehouse_id": seed_data["warehouse_id"]},
            headers=auth_headers,
        )

        assert resp.status_code == 200
        assert resp.json()["status"] == "RECEIVED"


class TestCancellations:
    def test_full_flow(self, client, auth_headers, api_session_factory, seed_data):
        option_id = _seed_product_option(api_session_factory)
        ctx = _create_order_with_item(api_session_factory, seed_data, "CANCEL-API-1", option_id)

        created = client.post(
            "/api/cancellations",
            json={"order_id": ctx["order_id"], "reason": "단순 변심", "refund_amount": 10000},
            headers=auth_headers,
        )
        assert created.status_code == 201
        cancellation_id = created.json()["id"]

        changed = client.patch(
            f"/api/cancellations/{cancellation_id}/status",
            json={"status": "COMPLETED", "warehouse_id": seed_data["warehouse_id"]},
            headers=auth_headers,
        )
        assert changed.status_code == 200
        assert changed.json()["status"] == "COMPLETED"

        order_resp = client.get(f"/api/orders/{ctx['order_id']}", headers=auth_headers)
        assert order_resp.json()["status"] == "CANCELED"

    def test_invalid_status_returns_422(self, client, auth_headers, api_session_factory, seed_data):
        option_id = _seed_product_option(api_session_factory)
        ctx = _create_order_with_item(api_session_factory, seed_data, "CANCEL-API-2", option_id)
        created = client.post(
            "/api/cancellations", json={"order_id": ctx["order_id"], "reason": "x"}, headers=auth_headers
        )

        resp = client.patch(
            f"/api/cancellations/{created.json()['id']}/status", json={"status": "APPROVED"}, headers=auth_headers
        )

        assert resp.status_code == 422


class TestExports:
    """SRS FR-REPORT-04: 교환/반품/취소 내역 엑셀 내보내기."""

    def test_export_exchanges_returns_xlsx_with_correct_content(
        self, client, auth_headers, api_session_factory, seed_data
    ):
        from io import BytesIO

        from openpyxl import load_workbook

        option_id = _seed_product_option(api_session_factory)
        ctx = _create_order_with_item(api_session_factory, seed_data, "EXPORT-EX-1", option_id)
        client.post(
            "/api/exchanges",
            json={"order_id": ctx["order_id"], "order_item_id": ctx["item_id"], "reason": "사이즈 교환"},
            headers=auth_headers,
        )

        resp = client.get("/api/exchanges/export", headers=auth_headers)

        assert resp.status_code == 200
        assert resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        wb = load_workbook(BytesIO(resp.content))
        ws = wb.active
        assert ws.title == "교환내역"
        assert ws.max_row >= 2

    def test_export_returns_returns_xlsx_with_correct_content(
        self, client, auth_headers, api_session_factory, seed_data
    ):
        from io import BytesIO

        from openpyxl import load_workbook

        option_id = _seed_product_option(api_session_factory, seed_data)
        ctx = _create_order_with_item(api_session_factory, seed_data, "EXPORT-RET-1", option_id)
        client.post(
            "/api/returns",
            json={
                "order_id": ctx["order_id"],
                "order_item_id": ctx["item_id"],
                "reason": "상품 불량",
                "refund_amount": 5000,
            },
            headers=auth_headers,
        )

        resp = client.get("/api/returns/export", headers=auth_headers)

        assert resp.status_code == 200
        wb = load_workbook(BytesIO(resp.content))
        ws = wb.active
        assert ws.title == "반품내역"
        assert ws.max_row >= 2

    def test_export_cancellations_returns_xlsx_with_correct_content(
        self, client, auth_headers, api_session_factory, seed_data
    ):
        from io import BytesIO

        from openpyxl import load_workbook

        option_id = _seed_product_option(api_session_factory)
        ctx = _create_order_with_item(api_session_factory, seed_data, "EXPORT-CANCEL-1", option_id)
        client.post(
            "/api/cancellations",
            json={"order_id": ctx["order_id"], "reason": "단순 변심", "refund_amount": 10000},
            headers=auth_headers,
        )

        resp = client.get("/api/cancellations/export", headers=auth_headers)

        assert resp.status_code == 200
        wb = load_workbook(BytesIO(resp.content))
        ws = wb.active
        assert ws.title == "취소내역"
        assert ws.max_row >= 2

    def test_export_requires_authentication(self, client, seed_data):
        assert client.get("/api/exchanges/export").status_code == 401
        assert client.get("/api/returns/export").status_code == 401
        assert client.get("/api/cancellations/export").status_code == 401


class TestRBAC:
    def test_requires_exchange_return_manage_permission(self, client, api_session_factory, seed_data):
        from core.security import hash_password
        from models.user import Permission, Role, RolePermission, User

        db = api_session_factory()
        try:
            role = Role(name="ViewOnly")
            db.add(role)
            db.flush()
            other_perm = db.query(Permission).filter_by(code="ORDER_VIEW").first()
            db.add(RolePermission(role_id=role.id, permission_id=other_perm.id))
            db.add(
                User(
                    username="viewonly2",
                    password_hash=hash_password("pw123456"),
                    name="조회전용",
                    role_id=role.id,
                    is_active=True,
                )
            )
            db.commit()
        finally:
            db.close()

        login = client.post("/api/auth/login", data={"username": "viewonly2", "password": "pw123456"})
        headers = {"Authorization": f"Bearer {login.json()['access_token']}"}

        assert client.get("/api/exchanges", headers=headers).status_code == 403
        assert client.get("/api/returns", headers=headers).status_code == 403
        assert client.get("/api/cancellations", headers=headers).status_code == 403
